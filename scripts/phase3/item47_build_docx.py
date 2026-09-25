#!/usr/bin/env python
"""Item 47. Convert MANUSCRIPT_KMD_mycobiome.md to a Word document for inspection.

Times New Roman 12 pt, double spaced, 2.5 cm margins, continuous line numbers, page numbers.
Markdown emphasis is carried into real Word italics, tables become native Word tables, figures
are embedded at their callout positions, and every remaining INSERT is bold and highlighted.
A figure that will not embed leaves a visible placeholder; nothing is dropped silently.
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

K = r"<KMD_ROOT>"
O = os.path.join(K, "__reanalysis_2026-06", "phase3", "outputs_27-07-2026")
CLEAN = os.path.join(O, "clean")
SRC = os.path.join(K, "MANUSCRIPT_KMD_mycobiome.md")
OUT = os.path.join(K, "MANUSCRIPT_KMD_mycobiome.docx")

FIGFILE = {
 "Figure 1": (os.path.join(O, "figures", "Figure1_host_by_site.png"),
              "Host DNA content by site."),
 "Figure 2": (os.path.join(CLEAN, "figures", "Figure2_bloem_hostload_vs_ASVs.png"),
              "Bloemfontein: apparent fungal richness against host load."),
 "Figure 3": (os.path.join(O, "figures", "Figure3_tremellomycetes_placement.png"),
              "Phylogenetic placement of the unresolved Tremellomycetes lineage."),
 "Figure 4": (os.path.join(CLEAN, "figures", "Figure4_composition.png"),
              "Taxonomic composition by site and condition."),
 "Figure 5": (os.path.join(CLEAN, "figures", "Figure5_hill_numbers.png"),
              "Hill numbers q0, q1 and q2 by site and condition, at depth 57,124."),
 "Figure 6": (os.path.join(CLEAN, "figures", "Figure6_ordination.png"),
              "Bray-Curtis and Jaccard ordination of the 16 core samples."),
 "Figure 7": (os.path.join(CLEAN, "figures", "Figure7_curvibasidium.png"),
              "Per-sample Curvibasidium relative abundance."),
 "Figure 8": (os.path.join(O, "figures", "Figure8_placement_by_neighbourhood.png"),
              "Placement of abundant unnamed ASVs, by taxonomic neighbourhood."),
 "Figure 9": (os.path.join(CLEAN, "figures", "Figure9_rare_fraction.png"),
              "The low-abundance fraction of the fungal community. Panel a, abundance distribution of all ASVs with the 1 percent threshold marked. Panel b, taxonomic resolution by abundance class."),
}
# figure is inserted after the paragraph that first cites it
FIRST_CITE = {}

INSERT_RE = re.compile(r"\[INSERT[^\]]*\]")
ITAL_RE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*")

def set_font(run, italic=None, bold=None, highlight=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    r = run._element.rPr.rFonts
    r.set(qn("w:eastAsia"), "Times New Roman")
    if italic is not None: run.italic = italic
    if bold is not None: run.bold = bold
    if highlight:
        h = OxmlElement("w:highlight"); h.set(qn("w:val"), "yellow")
        run._element.rPr.append(h)

def add_runs(par, text):
    """Split markdown emphasis and INSERT markers into correctly formatted runs."""
    pos = 0
    tokens = []
    for m in re.finditer(r"\[INSERT[^\]]*\]|\*\*(?:.+?)\*\*|\*(?:.+?)\*", text):
        if m.start() > pos: tokens.append(("plain", text[pos:m.start()]))
        s = m.group(0)
        if s.startswith("[INSERT"): tokens.append(("insert", s))
        elif s.startswith("**"): tokens.append(("bold", s[2:-2]))
        else: tokens.append(("ital", s[1:-1]))
        pos = m.end()
    if pos < len(text): tokens.append(("plain", text[pos:]))
    if not tokens: tokens = [("plain", text)]
    for kind, s in tokens:
        if not s: continue
        run = par.add_run(s)
        if kind == "ital": set_font(run, italic=True)
        elif kind == "bold": set_font(run, bold=True)
        elif kind == "insert": set_font(run, bold=True, highlight=True)
        else: set_font(run)

def style_par(p, space_after=6):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)

# continuous line numbers
sectPr = sec._sectPr
ln = OxmlElement("w:lnNumType")
ln.set(qn("w:countBy"), "1"); ln.set(qn("w:start"), "1")
ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "360")
sectPr.append(ln)

# page number in the footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run(); set_font(r)
for el, attr in (("w:fldChar", {"w:fldCharType": "begin"}),
                 ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "end"})):
    e = OxmlElement(el)
    if attr:
        for k, v in attr.items(): e.set(qn(k), v)
    else:
        e.set(qn("xml:space"), "preserve"); e.text = " PAGE "
    r._element.append(e)

for lvl, sz in ((1, 16), (2, 14), (3, 12)):
    s = doc.styles[f"Heading {lvl}"]
    s.font.name = "Times New Roman"; s.font.size = Pt(sz)
    s.font.bold = True; s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

lines = open(SRC, encoding="utf-8").read().split("\n")
missing_figs, inserted_figs, n_tables = [], [], 0
i = 0
def emit_figure(num):
    global missing_figs, inserted_figs
    if num in inserted_figs or num not in FIGFILE: return
    path, cap = FIGFILE[num]
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_par(p, space_after=2)
    if os.path.exists(path):
        try:
            p.add_run().add_picture(path, width=Cm(15.5))
            inserted_figs.append(num)
        except Exception as e:
            set_font(p.add_run(f"[FIGURE NOT EMBEDDED: {num}, {type(e).__name__}]"),
                     bold=True, highlight=True)
            missing_figs.append((num, f"{type(e).__name__}: {e}"))
    else:
        set_font(p.add_run(f"[FIGURE FILE NOT FOUND: {num} at {path}]"), bold=True, highlight=True)
        missing_figs.append((num, "file not found"))
    c = doc.add_paragraph(); style_par(c, space_after=12)
    cr = c.add_run(f"{num}. "); set_font(cr, bold=True)
    add_runs(c, cap)

while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1; continue
    if line.strip() == "---":
        i += 1; continue
    # table block
    if line.startswith("|"):
        # The drafts carry a blank line between table rows, a leftover from the docx
        # conversion. Stopping at the first blank line splits one table into a row of
        # one-row tables, so look past blank lines for a continuing row.
        block = []
        while i < len(lines):
            if lines[i].startswith("|"):
                block.append(lines[i].rstrip()); i += 1; continue
            if not lines[i].strip():
                j = i
                while j < len(lines) and not lines[j].strip(): j += 1
                if j < len(lines) and lines[j].startswith("|"):
                    i = j; continue
            break
        rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in block
                if not re.match(r"^\|[\s\-:|]+\|$", r.strip())]
        if rows:
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    if ci >= len(rows[0]): continue
                    cp = t.cell(ri, ci).paragraphs[0]
                    cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    cp.paragraph_format.space_after = Pt(0)
                    add_runs(cp, cell)
                    if ri == 0:
                        for rr in cp.runs: rr.bold = True
            n_tables += 1
            doc.add_paragraph()
        continue
    m = re.match(r"^(#{1,3})\s+(.*)$", line)
    if m:
        lvl = len(m.group(1))
        h = doc.add_heading(level=lvl)
        h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        txt = re.sub(r"\*+", "", m.group(2))
        r2 = h.add_run(txt)
        r2.font.name = "Times New Roman"; r2.font.bold = True
        r2.font.size = Pt({1: 16, 2: 14, 3: 12}[lvl])
        r2.font.color.rgb = RGBColor(0, 0, 0)
        i += 1; continue
    if line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet"); style_par(p)
        add_runs(p, line[2:])
        i += 1; continue
    # A markdown paragraph is a run of consecutive non-blank lines. The Introduction and
    # Discussion are hard-wrapped at about 80 characters, so treating each source line as its
    # own Word paragraph would render those sections as hundreds of one-line paragraphs.
    para = [line]
    k = i + 1
    while k < len(lines):
        nxt = lines[k].rstrip()
        if (not nxt.strip() or nxt.startswith(("|", "#", "- ")) or nxt.strip() == "---"):
            break
        para.append(nxt); k += 1
    line = " ".join(para); i = k - 1
    p = doc.add_paragraph(); style_par(p)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, line)
    # place each figure after the paragraph that first cites it
    for num in FIGFILE:
        if num not in inserted_figs and re.search(r"\b" + num.replace(" ", r"\s+") + r"\b", line) \
           and "Supplementary" not in line:
            emit_figure(num)
    i += 1

# any figure never cited in a body paragraph still has to appear
for num in FIGFILE:
    if num not in inserted_figs and not any(n == num for n, _ in missing_figs):
        doc.add_page_break(); emit_figure(num)

doc.save(OUT)
print(f"wrote {OUT}")
print(f"   figures embedded: {len(inserted_figs)} of {len(FIGFILE)}  {inserted_figs}")
print(f"   tables converted to native Word tables: {n_tables}")
if missing_figs:
    print("   FIGURES NOT EMBEDDED (placeholders left in the document):")
    for n, why in missing_figs: print(f"      {n}: {why}")
else:
    print("   no figure placeholders needed")
