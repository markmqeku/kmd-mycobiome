#!/usr/bin/env python
"""KMD PROMPT 4, R2g, R2h, R5. Per-library metadata for this study's libraries only (Table S12), the sampling
table (Table 1) and the number of Bloemfontein libraries in which the Tremellomycetes lineage was detected.
Sources: metadata_v3, ANSWERS.md revision 4 (pooling), mEI.xlsx (male and female trees, confirmed by Mark),
the Bioanalyzer transcription (r0a_bioanalyzer.json), the full and corrected feature tables, the host ASV list."""
import csv, json, re
from pathlib import Path
import openpyxl

ROOT = Path(r"<KMD_ROOT>")
O = ROOT / "__reanalysis_2026-06" / "phase3" / "outputs_27-07-2026"
meta = {r["sample-id"]: r for r in csv.DictReader(open(O / "clean" / "metadata_v3_27-07-2026.tsv", encoding="utf-8"), delimiter="\t")
        if not r["sample-id"].startswith("#")}


def table(p):
    rows = [l.rstrip("\n").split("\t") for l in open(p, encoding="utf-8") if not l.startswith("# ")]
    cols = rows[0][1:]
    return cols, {r[0]: [float(x) for x in r[1:]] for r in rows[1:]}


fc, full = table(O / "table_exp" / "feature-table.tsv")
cc, clean = table(O / "clean" / "feature-table-clean.tsv")
host = {l.split()[0] for l in open(O / "nontarget_screen" / "host_asv_ids_ALL.txt", encoding="utf-8") if l.strip()}
lineage = set(re.findall(r"md5=([0-9a-f]{32})", open(O / "clean" / "genbank_submission" / "Tremellomycetes_ASVs.fasta", encoding="utf-8").read()))
tot = {c: sum(v[i] for v in full.values()) for i, c in enumerate(fc)}
hst = {c: sum(full[a][i] for a in host if a in full) for i, c in enumerate(fc)}
fun = {c: sum(v[i] for v in clean.values()) for i, c in enumerate(cc)}
lin = {c: sum(clean[a][i] for a in lineage if a in clean) for i, c in enumerate(cc)}

sex = {}
ws = openpyxl.load_workbook(ROOT / "SOURCE_LAB_RECORDS" / "mEI.xlsx").worksheets[0]
for r in range(3, 35):
    m = re.search(r"\((M|F)\)", ws.cell(r, 2).value)
    sex[ws.cell(r, 1).value.strip()] = {"M": "male", "F": "female"}[m.group(1)] if m else ""
ba = {}
for w in json.load(open(ROOT / "__reanalysis_2026-06/phase3/r0a_bioanalyzer.json", encoding="utf-8")):
    if not re.fullmatch(r"P\d+", w["name"]): continue
    lib = [p for p in w["peaks"] if 300 <= p["size_bp"] <= 800 and not p["obs"]]
    val = (f"{max(lib, key=lambda p: p['conc_pg_ul'])['size_bp']:.0f} bp ({w['date'][8:]} Aug 2016)" if lib
           else f"not sized, failed markers ({w['date'][8:]} Aug 2016)")
    ba.setdefault(w["name"], []).append(val)

EXCL = {"P6": "excluded: sequence file truncated in delivery (no Bioanalyzer record)",
        "P15": "excluded: sequence file truncated in delivery; normal library peak on the Bioanalyzer",
        "P24": "excluded: sequence file truncated in delivery; normal library peak on the Bioanalyzer",
        "YT-P": "excluded: sequence file corrupted"}
TIS = {"Leaves": "leaf", "Twigs": "twig", "Inflorescence": "inflorescence", "Seeds": "seed", "Malformation": "malformed shoot"}
RUN = {"Bloemfontein": "run 3 (2016, 2 x 250 bp)", "Christiana": "run 13 (2017, 2 x 301 bp)", "Pretoria": "run 13 (2017, 2 x 301 bp)"}
rows = []
order = sorted(meta, key=lambda s: ({"Christiana": 0, "Pretoria": 1, "Bloemfontein": 2}[meta[s]["location"]],
                                     int(s[1:]) if re.fullmatch(r"P\d+", s) else 0, s))
for s in order:
    m = meta[s]; site = m["location"]
    cond = {"reference_site": "reference site"}.get(m["condition_std"], m["condition_std"])
    included = s not in EXCL
    status = "included" + ("; no fungal reads after host removal, absent from composition figures" if included and not fun.get(s) else "") if included else EXCL[s]
    rows.append({
        "library": s, "site": site, "condition": cond, "tissue": TIS[m["tissue"]],
        "age class": m["age"] if m["age"] and m["age"] != "not specified" else "not recorded",
        "trees pooled": "1" if site == "Bloemfontein" else "4",
        "tree sex": sex.get(s, "") or ("not recorded" if site == "Bloemfontein" else "not recorded"),
        "sequencing run": RUN[site],
        "Bioanalyzer main peak": "; ".join(ba.get(s, [])) or ("no record" if site == "Bloemfontein" else "no record (none exist for 2017)"),
        "total reads": f"{int(tot[s]):,}" if s in tot else "not denoised",
        "host percent": f"{100*hst[s]/tot[s]:.1f}" if tot.get(s) else "not applicable",
        "fungal reads": f"{int(fun.get(s, 0)):,}" if s in tot else "not applicable",
        "status": status})
rows.append({"library": "30-C", "site": "Pretoria", "condition": "asymptomatic", "tissue": "not recorded", "age class": "not recorded",
             "trees pooled": "4", "tree sex": "not recorded", "sequencing run": "not recorded", "Bioanalyzer main peak": "no record",
             "total reads": "not applicable", "host percent": "not applicable", "fungal reads": "not applicable",
             "status": "excluded: collected, but its sequence file was corrupted and it was never used"})
out = O / "clean" / "TableS12_library_metadata.tsv"
with open(out, "w", encoding="utf-8", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=list(rows[0]), delimiter="\t", lineterminator="\n"); w.writeheader(); w.writerows(rows)
print(f"Table S12: {len(rows)} libraries written to {out.name}")
n_inc = sum(1 for r in rows if r["status"].startswith("included"))
print("included:", n_inc, "| excluded:", [r["library"] for r in rows if not r["status"].startswith("included")])
bl = [c for c in cc if meta[c]["location"] == "Bloemfontein"]
det = [c for c in bl if lin[c] > 0]
print(f"R5: lineage detected in {len(det)} of 29 usable Bloemfontein libraries (P2 has no fungal reads); absent from:",
      sorted(set(c for c in bl if lin[c] == 0) | {"P2"}, key=lambda s: int(s[1:])))
core = [c for c in cc if meta[c]["location"] != "Bloemfontein"]
print(f"lineage detected in {sum(1 for c in core if lin[c] > 0)} of 16 core libraries; overall {len(det) + sum(1 for c in core if lin[c] > 0)} of 45")
