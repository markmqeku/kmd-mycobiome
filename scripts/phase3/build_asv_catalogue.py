#!/usr/bin/env python
"""
T-1 / T-2: ASV display labels + catalogue workbook.

Scientific rules enforced here:
  * The md5 ASV ID (sequence-derived) remains the PRIMARY KEY in every output.
    display_label is an ADDITIONAL human-readable column, never a replacement.
  * A name is only written at a rank the classifier actually returned.
    Empty ranks ("g__"), and the placeholder tokens UNITE uses for
    unresolved lineages ("unidentified", "incertae_sedis", "unclassified"),
    are treated as NOT achieved. Nothing is inferred below the achieved rank.
  * Numbering is by descending total read abundance across the 45 samples.
Usage: build_asv_catalogue.py <taxonomy.tsv> <rep.fasta> <table.tsv> <alignment.fasta|NONE> <outdir> <db_label>
"""
import sys, csv, hashlib, os, datetime
from collections import Counter, OrderedDict

RANKS = [("k__","kingdom"),("p__","phylum"),("c__","class"),("o__","order"),
         ("f__","family"),("g__","genus"),("s__","species")]
PLACEHOLDERS = {"unidentified","unclassified","incertae_sedis","unknown",""}

def parse_taxon(tax):
    """Return OrderedDict rank->name, only for ranks genuinely achieved."""
    out = OrderedDict((r[1], "") for r in RANKS)
    if not tax or tax.strip().lower().startswith("unassigned"):
        return out
    for part in tax.split(";"):
        part = part.strip()
        for pref, rank in RANKS:
            if part.startswith(pref):
                name = part[len(pref):].strip()
                base = name.lower().replace(" ", "_")
                # UNITE placeholder conventions -> rank NOT achieved:
                #   *_gen_Incertae_sedis / *_fam_Incertae_sedis  (rank unresolved)
                #   s__<HigherTaxon>_sp                          (species unresolved, e.g. s__Dothideales_sp)
                # A genuine UNITE species is binomial (Genus_epithet); a literal '_sp' epithet is a placeholder.
                if (base in PLACEHOLDERS
                        or base.endswith("_incertae_sedis")
                        or base.startswith("unidentified")
                        or (rank == "species" and (base.endswith("_sp") or base.endswith("_sp.")))):
                    name = ""
                out[rank] = name
    return out

def finest(ranks):
    """(rank_name, value) of the finest rank actually achieved, else (None, None)."""
    last = (None, None)
    for _, rank in RANKS:
        v = ranks.get(rank, "")
        if v:
            last = (rank, v)
    return last

def display_label(idx, ranks):
    """ASV####_<finest achieved name>, with explicit rank honesty."""
    tag = f"ASV{idx:04d}"
    rank, val = finest(ranks)
    if rank is None:
        return f"{tag}_unassigned"
    if rank == "species":
        sp = val
        g = ranks.get("genus", "")
        # UNITE species are usually already 'Genus_species'
        if g and not sp.lower().startswith(g.lower()):
            sp = f"{g}_{sp}"
        return f"{tag}_{sp}"
    if rank == "genus":
        return f"{tag}_{val}_sp"          # genus reached, species not: '_sp' (not invented)
    return f"{tag}_{val}_unassigned"      # finest achieved is above genus

def read_fasta(path):
    seqs, sid, buf = {}, None, []
    with open(path) as fh:
        for line in fh:
            line = line.rstrip("\n")
            if line.startswith(">"):
                if sid: seqs[sid] = "".join(buf)
                sid, buf = line[1:].split()[0], []
            else:
                buf.append(line.strip())
    if sid: seqs[sid] = "".join(buf)
    return seqs

def main():
    tax_p, fa_p, tab_p, aln_p, outdir, db_label = sys.argv[1:7]
    os.makedirs(outdir, exist_ok=True)

    # taxonomy
    tax, conf = {}, {}
    with open(tax_p) as fh:
        rd = csv.reader(fh, delimiter="\t"); hdr = next(rd)
        ci = hdr.index("Consensus") if "Consensus" in hdr else (hdr.index("Confidence") if "Confidence" in hdr else None)
        for row in rd:
            if not row or row[0].startswith("#"): continue
            tax[row[0]] = row[1]
            conf[row[0]] = row[ci] if ci is not None and ci < len(row) else ""

    seqs = read_fasta(fa_p)
    aln  = read_fasta(aln_p) if aln_p != "NONE" and os.path.exists(aln_p) else {}

    # feature table (biom-converted tsv: rows=ASV, cols=samples)
    with open(tab_p) as fh:
        lines = [l.rstrip("\n") for l in fh if l.strip() and not l.startswith("# Constructed")]
    hdr = lines[0].lstrip("#").split("\t")
    samples = [h for h in hdr[1:] if h.strip()]
    counts = {}
    for l in lines[1:]:
        p = l.split("\t")
        counts[p[0]] = [int(float(x)) for x in p[1:1+len(samples)]]

    asvs = list(counts.keys())
    total = {a: sum(counts[a]) for a in asvs}
    prev  = {a: sum(1 for v in counts[a] if v > 0) for a in asvs}
    asvs.sort(key=lambda a: (-total[a], a))          # descending abundance

    parsed = {a: parse_taxon(tax.get(a, "")) for a in asvs}
    label  = {a: display_label(i+1, parsed[a]) for i, a in enumerate(asvs)}

    # ---- reports ----
    rep = []
    rep.append(f"ASVs: {len(asvs)} | samples: {len(samples)} | total reads: {sum(total.values())}")
    rep.append(f"Database: {db_label}")
    rep.append("\nRANK RESOLUTION (n and % of ASVs with a genuine name at that rank):")
    for _, rank in RANKS:
        n = sum(1 for a in asvs if parsed[a][rank])
        rep.append(f"  {rank:8}: {n:5}  ({100*n/len(asvs):5.1f}%)")
    n_un = sum(1 for a in asvs if finest(parsed[a])[0] is None)
    rep.append(f"  {'none':8}: {n_un:5}  ({100*n_un/len(asvs):5.1f}%)  [fully unassigned]")

    rep.append("\nASV-per-TAXON COLLAPSE (multiple ASVs sharing one finest-rank name):")
    coll = Counter()
    for a in asvs:
        r, v = finest(parsed[a])
        coll[f"{v} ({r})" if r else "unassigned"] += 1
    multi = [(k, v) for k, v in coll.items() if v > 1]
    rep.append(f"  distinct taxon names: {len(coll)}; names carrying >1 ASV: {len(multi)}")
    for k, v in sorted(multi, key=lambda x: -x[1])[:25]:
        rep.append(f"    {k}: {v} ASVs")

    report = "\n".join(rep)
    print(report)
    open(os.path.join(outdir, "T1_label_report.txt"), "w", encoding="utf-8").write(report + "\n")

    # ---- FASTA + TSV interchange ----
    with open(os.path.join(outdir, "ASV_catalogue.fasta"), "w", encoding="utf-8") as fh:
        for a in asvs:
            fh.write(f">{label[a]}\n{seqs.get(a,'')}\n")
    with open(os.path.join(outdir, "ASV_catalogue.tsv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["ASV_ID_md5","display_label","sequence","length_bp","taxonomy_string",
                    *[r for _, r in RANKS],"classifier_confidence","total_reads","prevalence_n_samples"])
        for a in asvs:
            w.writerow([a,label[a],seqs.get(a,""),len(seqs.get(a,"")),tax.get(a,""),
                        *[parsed[a][r] for _, r in RANKS],conf.get(a,""),total[a],prev[a]])

    # ---- workbook ----
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "ASV_catalogue"
    ws.append(["ASV_ID_md5","display_label","representative_sequence","length_bp","taxonomy_string",
               *[r for _, r in RANKS],"classifier_confidence","total_reads","prevalence_n_samples"])
    for a in asvs:
        ws.append([a,label[a],seqs.get(a,""),len(seqs.get(a,"")),tax.get(a,""),
                   *[parsed[a][r] for _, r in RANKS],conf.get(a,""),total[a],prev[a]])

    ws2 = wb.create_sheet("ASV_by_sample")
    ws2.append(["ASV_ID_md5","display_label",*samples])
    for a in asvs:
        ws2.append([a,label[a],*counts[a]])

    ws3 = wb.create_sheet("Taxonomy_summary")
    ws3.append(["rank","n_ASVs_resolved","pct_ASVs_resolved"])
    for _, rank in RANKS:
        n = sum(1 for a in asvs if parsed[a][rank])
        ws3.append([rank, n, round(100*n/len(asvs), 1)])
    ws3.append([]); ws3.append(["TOP 30 TAXA BY TOTAL READS","finest_rank","n_ASVs","total_reads"])
    agg = {}
    for a in asvs:
        r, v = finest(parsed[a]); key = (v or "unassigned", r or "none")
        agg.setdefault(key, [0, 0]); agg[key][0] += 1; agg[key][1] += total[a]
    for (v, r), (n, tot) in sorted(agg.items(), key=lambda x: -x[1][1])[:30]:
        ws3.append([v, r, n, tot])
    ws3.append([]); ws3.append(["TOP 30 TAXA BY PREVALENCE","finest_rank","n_ASVs","summed_prevalence"])
    agg2 = {}
    for a in asvs:
        r, v = finest(parsed[a]); key = (v or "unassigned", r or "none")
        agg2.setdefault(key, [0, 0]); agg2[key][0] += 1; agg2[key][1] += prev[a]
    for (v, r), (n, pv) in sorted(agg2.items(), key=lambda x: -x[1][1])[:30]:
        ws3.append([v, r, n, pv])

    ws4 = wb.create_sheet("Provenance_README")
    prov = [
        ["KMD fungal ITS2 metabarcoding — ASV catalogue"],
        ["Generated", datetime.date.today().isoformat()],
        [],
        ["PRIMARY KEY", "ASV_ID_md5 (md5 of the representative sequence). display_label is additional, never a replacement."],
        ["Naming rule", "Names appear only at ranks the classifier actually returned. '_sp' = genus reached, species not resolved. '_unassigned' = finest achieved rank is above genus, or no assignment. Nothing inferred below the achieved rank."],
        [],
        ["PIPELINE"],
        ["Platform","Illumina MiSeq, ITS2, primers ITS3/ITS4 (White et al. 1990)"],
        ["QIIME 2","2024.10.1"],
        ["Primer removal","cutadapt trim-paired, --p-front-f GCATCGATGAAGAACGCAGC, --p-front-r TCCTCCGCTTATTGATATGC, --p-discard-untrimmed"],
        ["DADA2 group A (Bloemfontein, 2x250)","--p-trunc-len-f 228 --p-trunc-len-r 200"],
        ["DADA2 group B (Christiana + Pretoria, 2x301)","--p-trunc-len-f 280 --p-trunc-len-r 230"],
        ["Rationale","Sites were sequenced at different read lengths; a single truncation destroyed the 2x250 libraries in the original run. Denoised per read-length group, feature tables merged on md5 ASV IDs."],
        ["ITSxpress","Not used (excluded on measured retention grounds)"],
        ["Taxonomy",db_label],
        ["Classifier","classify-consensus-vsearch, --p-perc-identity 0.8 --p-maxaccepts 10"],
        ["Phylogeny","MAFFT -> mask -> FastTree -> midpoint root (ITS alignment is unreliable across distant fungi; phylogenetic metrics secondary)"],
        [],
        ["TOTALS"],
        ["Samples in table", len(samples)],
        ["ASVs", len(asvs)],
        ["Total reads", sum(total.values())],
        [],
        ["WRITE-OFFS (excluded, not in table)"],
        ["P6, P15, P24","raw files truncated to ~0 reads; no clean copy in tree"],
        ["YT-P","R1 gzip corrupt; R1/R2 shared only 834 read IDs (not a valid pair)"],
        ["OT-P","retained: one corrupt FASTQ record removed, re-paired to 110,302 matched pairs"],
        [],
        ["SAMPLES"], [", ".join(samples)],
    ]
    for row in prov: ws4.append(row)

    if aln:
        ws5 = wb.create_sheet("Aligned_sequences")
        ws5.append(["ASV_ID_md5","display_label","gapped_alignment_row"])
        ws5.append(["NOTE: MAFFT alignment used to build the tree. Distinct from the representative sequence in ASV_catalogue.","",""])
        for a in asvs:
            if a in aln: ws5.append([a, label[a], aln[a]])

    xlsx = os.path.join(outdir, "ASV_catalogue.xlsx")
    wb.save(xlsx)
    print(f"\nWROTE {xlsx}")
    print(f"WROTE {os.path.join(outdir,'ASV_catalogue.fasta')}")
    print(f"WROTE {os.path.join(outdir,'ASV_catalogue.tsv')}")

if __name__ == "__main__":
    main()
